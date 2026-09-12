from pathlib import Path
import sys
import docx
from docx.oxml.ns import qn

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import postprocess as pp


def field_instrs(doc):
    return [f.get(qn("w:instr"))
            for f in doc.element.body.iter(qn("w:fldSimple"))]


def texts(doc):
    return "\n".join(p.text for p in doc.paragraphs)


def test_build_toc_lists_emits_three_toc_fields(tmp_path, monkeypatch):
    monkeypatch.setattr(pp, "HERE", Path(__file__).resolve().parents[1])
    d = docx.Document()
    pp.build_toc_lists(d)
    instrs = field_instrs(d)
    assert any(i.startswith("TOC") and '\\o' in i for i in instrs)   # muc luc
    assert sum(i.startswith("TOC") for i in instrs) == 3
    body = texts(d)
    assert "MỤC LỤC" in body
    assert "DANH MỤC HÌNH" in body
    assert "DANH MỤC BẢNG" in body
    assert "DANH MỤC TỪ VIẾT TẮT" in body
    assert "ALPR" in body


def test_find_and_strip_before_abstract():
    d = docx.Document()
    d.add_paragraph("garbage title page")
    d.add_heading("Tóm tắt nội dung", level=1)
    d.add_paragraph("noi dung tom tat")
    idx = pp.find_abstract_index(d)
    assert idx >= 0
    pp.strip_before_abstract(d)
    assert "garbage title page" not in texts(d)
    assert "Tóm tắt nội dung" in texts(d)
