from pathlib import Path
import docx
from docx.shared import Pt, Cm

REF = Path(__file__).resolve().parents[1] / "reference.docx"


def test_normal_style_is_times_13pt():
    d = docx.Document(str(REF))
    normal = d.styles["Normal"]
    assert normal.font.name == "Times New Roman"
    assert normal.font.size == Pt(13)


def test_line_spacing_is_one_and_half():
    d = docx.Document(str(REF))
    assert abs(d.styles["Normal"].paragraph_format.line_spacing - 1.5) < 1e-6


def test_page_margins_match_bieumau():
    d = docx.Document(str(REF))
    sec = d.sections[0]
    assert abs(sec.top_margin - Cm(3)) < Cm(0.05)
    assert abs(sec.bottom_margin - Cm(3.5)) < Cm(0.05)
    assert abs(sec.left_margin - Cm(3.5)) < Cm(0.05)
    assert abs(sec.right_margin - Cm(2)) < Cm(0.05)
