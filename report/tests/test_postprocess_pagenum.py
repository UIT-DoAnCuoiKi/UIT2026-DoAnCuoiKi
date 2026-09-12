from pathlib import Path
import sys
import docx
from docx.oxml.ns import qn

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import postprocess as pp


def test_page_numbering_creates_second_section_starting_at_one():
    d = docx.Document()
    d.add_paragraph("front matter")
    abstract = d.add_heading("Tóm tắt nội dung", level=1)
    d.add_paragraph("body")

    pp.set_page_numbering(d, abstract)

    # Two sections now: front matter + body.
    assert len(d.sections) == 2
    body_sec = d.sections[1]
    sectPr = body_sec._sectPr
    pgNum = sectPr.find(qn("w:pgNumType"))
    assert pgNum is not None
    assert pgNum.get(qn("w:start")) == "1"
    # Body footer carries a PAGE field.
    footer_xml = body_sec.footer._element.xml
    assert "PAGE" in footer_xml
