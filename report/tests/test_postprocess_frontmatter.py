from pathlib import Path
import sys
import docx

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import postprocess as pp

META = {
    "title_vi": "He thong bai giu xe thong minh",
    "title_en": "Smart Parking System",
    "supervisor": "ThS. Phan Dinh Duy",
    "student_a": "Nguyen Minh Nhat", "student_a_id": "25410104",
    "student_b": "Le Quang Hoai Duc", "student_b_id": "25410034",
    "year": "2025-2026", "defense": "TP HCM, thang 9 nam 2026",
}


def texts(doc):
    return "\n".join(p.text for p in doc.paragraphs)


def test_cover_without_supervisor_omits_gvhd():
    d = docx.Document()
    pp.build_cover(d, META, with_supervisor=False)
    body = texts(d)
    assert "TRƯỜNG ĐẠI HỌC CÔNG NGHỆ THÔNG TIN" in body
    assert META["title_en"] in body
    assert "GIẢNG VIÊN HƯỚNG DẪN" not in body


def test_cover_with_supervisor_includes_gvhd_and_id():
    d = docx.Document()
    pp.build_cover(d, META, with_supervisor=True)
    body = texts(d)
    assert "GIẢNG VIÊN HƯỚNG DẪN" in body
    assert META["supervisor"] in body
    assert META["student_a_id"] in body


def test_council_page_has_heading():
    d = docx.Document()
    pp.build_council_page(d)
    assert "THÔNG TIN HỘI ĐỒNG CHẤM" in texts(d)


def test_acknowledgement_has_heading_and_text():
    d = docx.Document()
    pp.build_acknowledgement(d, "Xin cam on moi nguoi.")
    body = texts(d)
    assert "LỜI CẢM ƠN" in body
    assert "Xin cam on moi nguoi." in body
