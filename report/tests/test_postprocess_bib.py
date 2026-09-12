from pathlib import Path
import sys
import docx

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import postprocess as pp


def texts(doc):
    return [p.text for p in doc.paragraphs]


def test_split_bibliography_groups_by_lang_marker_and_strips_it():
    # citeproc (ieee-vi-en.csl) đã xếp Việt trước Anh và gắn marker {{lang:...}}.
    d = docx.Document()
    for label in ["{{lang:vi-VN}}[1] Nguyen VN paper",
                  "{{lang:vi-VN}}[2] Tran VN paper",
                  "{{lang:en-US}}[3] Smith EN paper"]:
        p = d.add_paragraph(label)
        p.style = d.styles["Normal"]
    pp.split_bibliography(d)
    body = texts(d)
    assert "Tài liệu tiếng Việt" in body
    assert "Tài liệu tiếng Anh" in body
    # Marker is stripped from every reference.
    assert not any("{{lang:" in t for t in body)
    # English heading sits before the (now unmarked) English reference, after the VN ones.
    en_ref = next(t for t in body if "Smith EN paper" in t)
    tran = next(t for t in body if "Tran VN paper" in t)
    assert body.index("Tài liệu tiếng Anh") < body.index(en_ref)
    assert body.index(tran) < body.index("Tài liệu tiếng Anh")


def test_split_bibliography_homogeneous_english_empty_vietnamese():
    d = docx.Document()
    for label in ["{{lang:en-US}}[1] Smith EN", "{{lang:en-US}}[2] Jones EN"]:
        p = d.add_paragraph(label)
        p.style = d.styles["Normal"]
    pp.split_bibliography(d)
    body = texts(d)
    # Both headings present; Anh heading immediately follows the empty Viet heading.
    assert body.index("Tài liệu tiếng Việt") < body.index("Tài liệu tiếng Anh")
    assert body.index("Tài liệu tiếng Anh") < body.index(next(t for t in body if "Smith EN" in t))
    assert not any("{{lang:" in t for t in body)


def test_uppercase_chapter_titles_skips_abstract():
    d = docx.Document()
    a = d.add_heading("Tóm tắt nội dung", level=1)
    c = d.add_heading("Giới thiệu", level=1)
    pp.uppercase_chapter_titles(d)
    assert "Tóm tắt nội dung" in [p.text for p in d.paragraphs]  # unchanged
    assert "GIỚI THIỆU" in [p.text for p in d.paragraphs]
